"""Applies the machine-readable fixes attached to issues. Works on live python-pptx shapes,
converting absolute slide coordinates back into each shape's own (possibly group-local) space."""
from pptx.util import Pt, Emu

from .config import Config
from .textfit import measure, overflow_bottom

EMU = 12700.0


def _pin(shape):
    """Write the full inherited geometry before moving; python-pptx otherwise leaves
    ext=0 on placeholders that inherit position from the layout and the box collapses."""
    l, t, w, h = shape.left, shape.top, shape.width, shape.height
    shape.left, shape.top, shape.width, shape.height = l, t, w, h


def _set_size(shape, size):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)


def set_abs(el, left=None, top=None, width=None, height=None):
    """Set absolute geometry (points) on el.shape, undoing any group transform."""
    sx, ox, sy, oy = el.xf
    sh = el.shape
    _pin(sh)
    if left is not None:
        sh.left = Emu(int(round((left * EMU - ox) / sx)))
    if top is not None:
        sh.top = Emu(int(round((top * EMU - oy) / sy)))
    if width is not None:
        sh.width = Emu(int(round(width * EMU / sx)))
    if height is not None:
        sh.height = Emu(int(round(height * EMU / sy)))
    if left is not None:
        el.left = left
    if top is not None:
        el.top = top
    if width is not None:
        el.width = width
    if height is not None:
        el.height = height


def _contained(outer, inner, pad=1.0):
    return outer.left - pad <= inner.left and outer.top - pad <= inner.top and \
        outer.right + pad >= inner.right and outer.bottom + pad >= inner.bottom


def move(sl, e, protect=(), **geom):
    """set_abs, but a filled panel/card takes the shapes sitting on it (same container) along."""
    old_l, old_t = e.left, e.top
    kids = []
    if e.kind == "shape" and e.fill is not None and not e.text.strip() and ("left" in geom or "top" in geom):
        kids = [o for o in sl.all if o.parent == e.parent and o is not e and o.id not in protect
                and o.kind not in ("group", "line") and _contained(e, o)]
    set_abs(e, **geom)
    dx, dy = e.left - old_l, e.top - old_t
    for o in kids:
        set_abs(o, left=o.left + dx if dx else None, top=o.top + dy if dy else None)


def refit_groups(shapes):
    """Make every group's frame hug its (moved) children again, keeping the group's scale so nothing shifts."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for sh in shapes:
        if sh.shape_type != 6:  # not a group
            continue
        refit_groups(sh.shapes)
        kids = [k for k in sh.shapes if None not in (k.left, k.top, k.width, k.height)]
        xfrm = sh._element.xpath("./p:grpSpPr/a:xfrm")
        if not kids or not xfrm:
            continue
        off, ext, cho, che = (xfrm[0].find(A + t) for t in ("off", "ext", "chOff", "chExt"))
        if None in (off, ext, cho, che):
            continue
        l = min(k.left for k in kids); t = min(k.top for k in kids)
        r = max(k.left + k.width for k in kids); b = max(k.top + k.height for k in kids)
        sx = float(ext.get("cx")) / (float(che.get("cx")) or 1.0)
        sy = float(ext.get("cy")) / (float(che.get("cy")) or 1.0)
        ox, oy = float(off.get("x")), float(off.get("y"))
        cx0, cy0 = float(cho.get("x")), float(cho.get("y"))
        off.set("x", str(int(round(ox + sx * (l - cx0))))); off.set("y", str(int(round(oy + sy * (t - cy0)))))
        ext.set("cx", str(int(round(sx * (r - l))))); ext.set("cy", str(int(round(sy * (b - t)))))
        cho.set("x", str(int(l))); cho.set("y", str(int(t)))
        che.set("cx", str(int(r - l))); che.set("cy", str(int(b - t)))


def apply(deck, issues, cfg: Config):
    """Mutates deck.prs; returns number of fixes applied. Marks issues fixed."""
    n = 0
    moved = set()  # (slide, shape_id, axis) already positioned
    order = {"remove": 0, "font": 1, "title_size": 1, "min_font": 2, "fit": 3, "title_fit": 3,
             "title_pos": 4, "align": 5, "place": 5, "resize": 5, "clamp": 3, "distribute": 6, "straighten": 7, "add_number": 8}
    todo = sorted((i for i in issues if i.fixable and i.fix), key=lambda i: order.get(i.fix["op"], 9))
    for iss in todo:
        sl = deck.slides[iss.slide - 1]
        els = {e.id: e for e in sl.all}
        f = iss.fix
        op = f["op"]
        ok = False
        if op == "remove":
            e = els.get(f["id"])
            if e is not None:
                e.shape._element.getparent().remove(e.shape._element)
                ok = True
        elif op == "font":
            e = els.get(f["id"])
            if e is not None:
                for p in e.shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            r.font.name = f["font"]
                ok = True
        elif op == "title_size":
            e = els.get(f["id"])
            if e is not None:
                _set_size(e.shape, f["size"])
                ok = True
        elif op == "min_font":
            e = els.get(f["id"])
            if e is not None:
                for p in e.shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt < cfg.min_font:
                            r.font.size = Pt(cfg.min_font)
                need = overflow_bottom(e, cfg, max(e.size or 0, cfg.min_font)) if e.kind == "text" else 0
                if need > e.height + 1 and e.top + need + 1 <= deck.height:
                    set_abs(e, height=need + 1)
                ok = True
        elif op in ("fit", "title_fit"):
            e = els.get(f["id"])
            if e is not None:
                floor = cfg.title_min_font if op == "title_fit" else cfg.min_font
                size = e.size or cfg.default_font

                m0 = measure(e, cfg, size)
                if m0["split"]:   # a word is wider than the box: widen it (keeps the type size)
                    from .textfit import text_width
                    word = max((text_width(w, size, e.bold) for w in e.text.split()), default=0)
                    need = word + m0["ml"] + m0["mr"] + 2
                    if e.left + need <= deck.width:
                        set_abs(e, width=max(e.width, need))

                def good(sz):
                    m = measure(e, cfg, sz)
                    if op == "title_fit":
                        return m["lines"] <= cfg.title_max_lines
                    return overflow_bottom(e, cfg, sz) <= e.height + 1 and not m["split"]
                while size > floor and not good(size):
                    size -= 1
                if good(size):
                    _set_size(e.shape, size)
                    ok = True
                elif op == "fit":  # last resort: grow the box downward if there is room
                    new_h = measure(e, cfg, max(size, floor))["height"]
                    if e.top + new_h <= deck.height:
                        _set_size(e.shape, max(size, floor))
                        set_abs(e, height=new_h)
                        ok = True
        elif op == "title_pos":
            e = els.get(f["id"])
            if e is not None:
                set_abs(e, left=f["pos"][0], top=f["pos"][1])
                ok = True
        elif op == "resize":
            e = els.get(f["id"])
            if e is not None:
                set_abs(e, width=f.get("width"), height=f.get("height"))
                ok = True
        elif op == "clamp":
            e = els.get(f["id"])
            if e is not None:
                l, r = max(e.left, f["l"]), min(e.right, f["r"])
                if r > l:
                    set_abs(e, left=l, width=r - l)
                    ok = True
        elif op == "place":
            e = els.get(f["id"])
            if e is not None:
                set_abs(e, left=f.get("left"), top=f.get("top"))
                ok = True
        elif op == "align":
            edge, target = f["edge"], f["target"]
            axis = "x" if edge in ("left", "right", "center") else "y"
            for sid in f["ids"]:
                e = els.get(sid)
                if e is None or (iss.slide, sid, axis) in moved:
                    continue
                prot = set(f["ids"])
                if edge == "left":
                    move(sl, e, prot, left=target)
                elif edge == "right":
                    move(sl, e, prot, left=target - e.width)
                elif edge == "center":
                    move(sl, e, prot, left=target - e.width / 2)
                elif edge == "top":
                    move(sl, e, prot, top=target)
                elif edge == "bottom":
                    move(sl, e, prot, top=target - e.height)
                moved.add((iss.slide, sid, axis))
            ok = True
        elif op == "distribute":
            grp = [els[i] for i in f["ids"] if i in els]
            if len(grp) >= 3:
                row = f["axis"] == "row"
                grp.sort(key=lambda s: s.left if row else s.top)
                size = (lambda s: s.width) if row else (lambda s: s.height)
                start = grp[0].left if row else grp[0].top
                end = (grp[-1].right if row else grp[-1].bottom)
                gap = (end - start - sum(size(s) for s in grp)) / (len(grp) - 1)
                pos = start
                for s in grp:
                    move(sl, s, set(f["ids"]), **({"left": pos} if row else {"top": pos}))
                    pos += size(s) + gap
                ok = True
        elif op == "straighten":
            e = els.get(f["id"])
            sh = e.shape if e is not None else None
            if sh is not None and hasattr(sh, "begin_x"):
                dx, dy = abs(sh.end_x - sh.begin_x), abs(sh.end_y - sh.begin_y)
                if dx >= dy:
                    sh.end_y = sh.begin_y
                else:
                    sh.end_x = sh.begin_x
                ok = True
        elif op == "add_number":
            l, t, w, h = f["ref"]
            tb = sl.slide.shapes.add_textbox(Emu(int(l * EMU)), Emu(int(t * EMU)), Emu(int(w * EMU)), Emu(int(h * EMU)))
            tb.text_frame.text = str(iss.slide)
            for p in tb.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(max(cfg.min_font, 12))
            ok = True
        if ok:
            iss.status = "fixed"
            n += 1
    for sl in deck.slides:
        refit_groups(sl.slide.shapes)
    return n
