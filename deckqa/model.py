"""Flat, unit-normalised (points) view of a PPTX deck, including shapes nested in groups."""
import re
from dataclasses import dataclass, field
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

EMU = 12700.0
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
IDENT = (1.0, 0.0, 1.0, 0.0)   # sx, ox, sy, oy : absolute_emu = o + s * local_emu


@dataclass
class El:
    shape: object
    slide: int
    id: int
    name: str
    left: float
    top: float
    width: float
    height: float
    kind: str                      # text | picture | line | shape | group | other
    parent: Optional[int] = None   # shape id of containing group (None = slide level)
    xf: tuple = IDENT              # local -> absolute transform of the shape's own coordinate space
    is_title: bool = False
    is_placeholder: bool = False
    is_number: bool = False
    ph_type: object = None
    text: str = ""
    sizes: list = field(default_factory=list)
    fonts: list = field(default_factory=list)
    implicit_runs: int = 0         # runs that inherit the theme font
    bold: bool = False
    fill: Optional[tuple] = None
    line: Optional[tuple] = None   # (x1, y1, x2, y2) absolute points for connectors
    has_pic: bool = False          # group that contains pictures
    n_leaves: int = 0              # for groups: number of leaf shapes inside
    group_has_text: bool = False

    @property
    def in_group(self): return self.parent is not None
    @property
    def right(self): return self.left + self.width
    @property
    def bottom(self): return self.top + self.height
    @property
    def cx(self): return self.left + self.width / 2
    @property
    def cy(self): return self.top + self.height / 2
    @property
    def bbox(self): return (self.left, self.top, self.right, self.bottom)

    @property
    def size(self):
        return max(self.sizes) if self.sizes else None


def _fill(shape):
    try:
        if shape.fill.type == 1:  # solid
            c = shape.fill.fore_color
            if c.type is not None and hasattr(c, "rgb") and c.rgb is not None:
                return tuple(c.rgb)
    except Exception:
        pass
    return None


def _runs(shape):
    if not getattr(shape, "has_text_frame", False):
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            yield r


def _abs(xf, l, t, w, h):
    sx, ox, sy, oy = xf
    return (ox + sx * l) / EMU, (oy + sy * t) / EMU, sx * w / EMU, sy * h / EMU


def load_element(shape, slide_no, cfg, slide_h, xf=IDENT, parent=None):
    try:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if None in (l, t, w, h):
        return None
    st = shape.shape_type
    is_ph = shape.is_placeholder
    ph_type = shape.placeholder_format.type if is_ph else None
    kind = "other"
    if st == MSO_SHAPE_TYPE.LINE or shape.__class__.__name__ == "Connector":
        kind = "line"
    elif st == MSO_SHAPE_TYPE.GROUP:
        kind = "group"
    elif st == MSO_SHAPE_TYPE.PICTURE or ph_type == PP_PLACEHOLDER.PICTURE:
        kind = "picture"
    elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        kind = "text"
    elif st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER):
        kind = "shape"
    al, at, aw, ah = _abs(xf, l, t, w, h)
    el = El(shape, slide_no, shape.shape_id, shape.name, al, at, aw, ah, kind, parent, xf)
    el.is_placeholder = is_ph
    el.ph_type = ph_type
    el.is_title = is_ph and ph_type in TITLE_TYPES
    text = shape.text_frame.text if getattr(shape, "has_text_frame", False) else ""
    el.text = text
    el.is_number = (is_ph and ph_type == PP_PLACEHOLDER.SLIDE_NUMBER) or (
        bool(re.fullmatch(r"\s*\d{1,3}\s*", text)) and el.top > slide_h * 0.85)
    for r in _runs(shape):
        if r.font.size:
            el.sizes.append(r.font.size.pt)
        if r.font.name:
            el.fonts.append(r.font.name)
        elif r.text.strip():
            el.implicit_runs += 1
        if r.font.bold:
            el.bold = True
    if el.kind == "text" and not el.sizes:
        el.sizes = [cfg.default_title_font if el.is_title else cfg.default_font]
    el.fill = _fill(shape) if kind in ("shape", "text") else None
    if kind == "line":
        try:
            bx, by, ex, ey = shape.begin_x, shape.begin_y, shape.end_x, shape.end_y
            sx, ox, sy, oy = xf
            el.line = ((ox + sx * bx) / EMU, (oy + sy * by) / EMU, (ox + sx * ex) / EMU, (oy + sy * ey) / EMU)
        except Exception:
            el.line = (el.left, el.top, el.right, el.bottom)
    return el


def _group_xf(group, xf):
    xfrm = group._element.xpath("./p:grpSpPr/a:xfrm")
    if not xfrm:
        return xf
    x = xfrm[0]
    def g(tag, attr):
        n = x.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag)
        return float(n.get(attr)) if n is not None else 0.0
    off_x, off_y = g("off", "x"), g("off", "y")
    ext_x, ext_y = g("ext", "cx"), g("ext", "cy")
    ch_x, ch_y = g("chOff", "x"), g("chOff", "y")
    che_x, che_y = g("chExt", "cx") or ext_x or 1.0, g("chExt", "cy") or ext_y or 1.0
    sx, ox, sy, oy = xf
    nsx = sx * ext_x / che_x if ext_x else sx
    nsy = sy * ext_y / che_y if ext_y else sy
    return (nsx, ox + sx * off_x - nsx * ch_x, nsy, oy + sy * off_y - nsy * ch_y)


def _walk(shapes, slide_no, cfg, H, xf, parent, out):
    for sh in shapes:
        el = load_element(sh, slide_no, cfg, H, xf, parent)
        if el is None:
            continue
        out.append(el)
        if el.kind == "group":
            before = len(out)
            _walk(sh.shapes, slide_no, cfg, H, _group_xf(sh, xf), sh.shape_id, out)
            kids = [e for e in out[before:] if e.kind != "group"]
            el.n_leaves = len(kids)
            el.has_pic = any(k.kind == "picture" for k in kids)
            el.group_has_text = any(k.kind == "text" for k in kids)


@dataclass
class Slide:
    no: int
    slide: object
    all: list                       # every element incl. groups, document order
    layout: str = ""

    @property
    def els(self):                  # slide-level elements (groups count as one box)
        return [e for e in self.all if e.parent is None]

    @property
    def leaves(self):               # everything that is not a group container
        return [e for e in self.all if e.kind != "group"]

    def siblings(self):             # parent id -> direct children
        d = {}
        for e in self.all:
            d.setdefault(e.parent, []).append(e)
        return d


@dataclass
class Deck:
    prs: object
    width: float
    height: float
    slides: list
    theme_fonts: tuple = ()      # (major, minor) latin typefaces from the theme


def _theme_fonts(prs):
    try:
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        master = prs.slide_masters[0]
        blob = master.part.part_related_by(RT.THEME).blob
        root = etree.fromstring(blob)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        return tuple(root.xpath(f"//a:{k}Font/a:latin/@typeface", namespaces=ns)[0] for k in ("major", "minor"))
    except Exception:
        return ()


def load(path, cfg):
    return build(Presentation(path), cfg)


def _guess_title(sl, H):
    """No title placeholder (common in consulting decks): biggest text near the top is the title."""
    if any(e.is_title for e in sl.all):
        return
    cands = [e for e in sl.leaves if e.kind == "text" and not e.is_number and e.top < H * 0.25
             and len(e.text.split()) >= 2]
    if cands:
        best = max(cands, key=lambda e: (round(e.size or 0), -e.top))
        best.is_title = True


def build(prs, cfg):
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    slides = []
    for i, s in enumerate(prs.slides, 1):
        out = []
        _walk(s.shapes, i, cfg, H, IDENT, None, out)
        sl = Slide(i, s, out, s.slide_layout.name)
        _guess_title(sl, H)
        slides.append(sl)
    return Deck(prs, W, H, slides, _theme_fonts(prs))
