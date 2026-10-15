import math
from collections import Counter, defaultdict
from statistics import median

from pptx.enum.text import MSO_AUTO_SIZE

from .issues import Issue
from .textfit import measure


def _bbox_union(els):
    return (min(e.left for e in els), min(e.top for e in els),
            max(e.right for e in els), max(e.bottom for e in els))


def _content(sl, deck):
    """Shapes that take part in layout checks (no lines, no full-slide backgrounds, no numbers)."""
    return [e for e in sl.els if e.kind in ("text", "picture", "shape", "group")
            and not e.is_number
            and not (e.width > deck.width * 0.95 and e.height > deck.height * 0.95)]


def _clusters(vals, tol):
    """vals: list of (value, el). Group values within tol of a neighbour."""
    vals = sorted(vals, key=lambda v: v[0])
    out, cur = [], [vals[0]] if vals else []
    for v in vals[1:]:
        if v[0] - cur[-1][0] <= tol:
            cur.append(v)
        else:
            out.append(cur)
            cur = [v]
    if cur:
        out.append(cur)
    return out


# ---------------------------------------------------------------- 1 Alignment
def alignment(deck, cfg):
    out = []
    edges = {"left": lambda e: e.left, "right": lambda e: e.right, "top": lambda e: e.top,
             "bottom": lambda e: e.bottom, "center": lambda e: e.cx}
    for sl in deck.slides:
        els = [e for e in _content(sl, deck) if not e.is_title]
        seen = set()  # same shapes already reported on another edge -> don't repeat
        for edge, fn in edges.items():
            if len(els) < 2:
                continue
            for cl in _clusters([(fn(e), e) for e in els], cfg.align_tol):
                vs = [v for v, _ in cl]
                if len(cl) < 2 or max(vs) - min(vs) <= cfg.align_eps:
                    continue
                if len({id(e) for _, e in cl}) < 2:
                    continue
                target = median(vs)
                members = [e for _, e in cl]
                key = frozenset(e.id for e in members)
                if key in seen:
                    continue
                seen.add(key)
                out.append(Issue(
                    sl.no, "Alignment",
                    f"{edge} edges of {len(members)} shapes are off by up to "
                    f"{max(vs) - min(vs):.1f}pt (should share one {edge} line).",
                    [e.id for e in members], _bbox_union(members), True,
                    dict(op="align", edge=edge, target=target, ids=[e.id for e in members])))
    return out


# ----------------------------------------------------------------- 2 Spacing
def spacing(deck, cfg):
    out = []
    for sl in deck.slides:
        els = [e for e in _content(sl, deck) if not e.is_title]
        # uneven gaps in rows and columns
        for axis in ("row", "col"):
            used = set()
            for e in els:
                if id(e) in used:
                    continue
                grp = [o for o in els if (
                    (min(e.bottom, o.bottom) - max(e.top, o.top) > min(e.height, o.height) * 0.5) if axis == "row"
                    else (min(e.right, o.right) - max(e.left, o.left) > min(e.width, o.width) * 0.5))]
                for g in grp:
                    used.add(id(g))
                if len(grp) < 3:
                    continue
                grp.sort(key=lambda s: s.left if axis == "row" else s.top)
                gaps = [(grp[i + 1].left - grp[i].right) if axis == "row" else (grp[i + 1].top - grp[i].bottom)
                        for i in range(len(grp) - 1)]
                if any(g < -0.5 for g in gaps):
                    continue  # overlapping layout, not a spacing row
                if max(gaps) - min(gaps) > cfg.gap_tol:
                    out.append(Issue(
                        sl.no, "Spacing",
                        f"Uneven gaps in a {axis} of {len(grp)} shapes "
                        f"({', '.join(f'{g:.0f}' for g in gaps)}pt).",
                        [s.id for s in grp], _bbox_union(grp), True,
                        dict(op="distribute", axis=axis, ids=[s.id for s in grp])))
        # empty band
        if els:
            srt = sorted(els, key=lambda s: s.top)
            reach = srt[0].bottom
            for s in srt[1:]:
                if s.top - reach > deck.height * cfg.band_ratio:
                    out.append(Issue(sl.no, "Spacing",
                                     f"Empty horizontal band of {s.top - reach:.0f}pt between content blocks.",
                                     [s.id], (0, reach, deck.width, s.top)))
                reach = max(reach, s.bottom)
    return out


# ---------------------------------------------------------------- 3 Text fit
def text_fit(deck, cfg):
    out = []
    for sl in deck.slides:
        for e in sl.els:
            if e.kind != "text" or e.is_number or e.is_title:
                continue  # titles are judged by the Titles check
            tf = e.shape.text_frame
            m = measure(e, cfg)
            if tf.word_wrap is not False and m["split"]:
                out.append(Issue(sl.no, "Text fit", "A word is wider than its box and will split mid-word.",
                                 [e.id], e.bbox, True, dict(op="fit", id=e.id)))
            auto = tf.auto_size in (MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)
            if m["height"] > e.height + 1.0 and not auto and e.height >= 1.0:  # height 0 = auto-growing box
                out.append(Issue(sl.no, "Text fit",
                                 f"Text needs ~{m['height']:.0f}pt but box is {e.height:.0f}pt tall (spills out).",
                                 [e.id], e.bbox, True, dict(op="fit", id=e.id)))
            elif m["orphan"]:
                out.append(Issue(sl.no, "Text fit", "Last line has a single hanging word.",
                                 [e.id], e.bbox, False))
    return out


# --------------------------------------------------------- 4 Lines and arrows
def _seg_hits_rect(x1, y1, x2, y2, r):
    l, t, rr, b = r
    # Liang-Barsky
    dx, dy = x2 - x1, y2 - y1
    p = [-dx, dx, -dy, dy]
    q = [x1 - l, rr - x1, y1 - t, b - y1]
    u1, u2 = 0.0, 1.0
    for pi, qi in zip(p, q):
        if pi == 0:
            if qi < 0:
                return False
        else:
            t_ = qi / pi
            if pi < 0:
                u1 = max(u1, t_)
            else:
                u2 = min(u2, t_)
    return u1 < u2 and (u2 - u1) > 0.02


def _inside(x, y, r, pad):
    return r[0] - pad <= x <= r[2] + pad and r[1] - pad <= y <= r[3] + pad


def lines_arrows(deck, cfg):
    out = []
    for sl in deck.slides:
        texts = [e for e in sl.els if e.kind == "text" and not e.is_number]
        for ln in (e for e in sl.els if e.kind == "line"):
            x1, y1, x2, y2 = ln.line
            dx, dy = x2 - x1, y2 - y1
            length = math.hypot(dx, dy)
            if length < 1:
                continue
            ang = math.degrees(math.atan2(abs(dy), abs(dx)))  # 0 horiz .. 90 vert
            off = min(ang, 90 - ang)
            if 0.2 < off <= cfg.angle_tol:
                out.append(Issue(sl.no, "Lines and arrows", f"Line is {off:.1f} degrees off the horizontal/vertical (crooked).",
                                 [ln.id], ln.bbox, True, dict(op="straighten", id=ln.id)))
            for t in texts:
                r = (t.left + 2, t.top + 2, t.right - 2, t.bottom - 2)
                if _inside(x1, y1, t.bbox, 2) or _inside(x2, y2, t.bbox, 2):
                    continue  # attached to the box, not crossing it
                if _seg_hits_rect(x1, y1, x2, y2, r):
                    out.append(Issue(sl.no, "Lines and arrows", "A line crosses through a text box.",
                                     [ln.id, t.id], _bbox_union([ln, t])))
    return out


# -------------------------------------------------------------------- 5 Titles
def titles(deck, cfg):
    out = []
    for sl in deck.slides:
        for t in (e for e in sl.els if e.is_title and e.kind == "text"):
            m = measure(t, cfg)
            if m["lines"] > cfg.title_max_lines:
                out.append(Issue(sl.no, "Titles", f"Title wraps to ~{m['lines']} lines (max {cfg.title_max_lines}).",
                                 [t.id], t.bbox, True, dict(op="title_fit", id=t.id)))
            text_bottom = t.top + m["height"]
            for o in _content(sl, deck):
                if o is t or o.is_title:
                    continue
                horiz = min(t.right, o.right) - max(t.left, o.left)
                if horiz > 0 and o.top < text_bottom - 0.5 and o.bottom > t.top + 0.5:
                    out.append(Issue(sl.no, "Titles", f"Title collides with '{o.name}'.",
                                     [t.id, o.id], _bbox_union([t, o])))
    return out


# --------------------------------------------------------------- 6 Readability
def readability(deck, cfg):
    out = []
    for sl in deck.slides:
        for e in sl.els:
            if e.kind != "text":
                continue
            small = [s for s in e.sizes if s < cfg.min_font]
            if small:
                out.append(Issue(sl.no, "Readability",
                                 f"Text is {min(small):g}pt (minimum {cfg.min_font:g}pt).",
                                 [e.id], e.bbox, True, dict(op="min_font", id=e.id)))
    return out


# ------------------------------------------------------------ 7 Leftover graphics
def _grey_or_black(rgb):
    if rgb is None:
        return False
    r, g, b = rgb
    return max(r, g, b) - min(r, g, b) <= 14


def leftovers(deck, cfg):
    out = []
    for sl in deck.slides:
        for e in sl.els:
            if e.kind == "picture" or e.kind == "line" or e.kind == "group":
                continue
            if e.left + e.width < 0 or e.top + e.height < 0 or e.left > deck.width or e.top > deck.height:
                out.append(Issue(sl.no, "Leftover graphics", f"'{e.name}' sits completely off the slide.",
                                 [e.id], e.bbox, True, dict(op="remove", id=e.id)))
                continue
            if e.kind == "shape" and not e.text.strip():
                if e.is_placeholder:
                    out.append(Issue(sl.no, "Leftover graphics", f"Empty placeholder '{e.name}'.",
                                     [e.id], e.bbox, True, dict(op="remove", id=e.id)))
                elif min(e.width, e.height) >= 6 and e.fill is not None:
                    why = "grey/black box" if _grey_or_black(e.fill) else "coloured blob"
                    out.append(Issue(sl.no, "Leftover graphics",
                                     f"Empty {why} '{e.name}' with no text or content.",
                                     [e.id], e.bbox, True, dict(op="remove", id=e.id)))
            elif e.kind == "text" and not e.text.strip():
                pass
    return out


# -------------------------------------------------------------- 8 Consistency
def consistency(deck, cfg):
    out = []
    # fonts
    fonts = Counter(f for sl in deck.slides for e in sl.els for f in e.fonts)
    if fonts:
        dom = fonts.most_common(1)[0][0]
        for sl in deck.slides:
            for e in sl.els:
                odd = sorted({f for f in e.fonts if f != dom})
                if odd:
                    out.append(Issue(sl.no, "Consistency",
                                     f"Font {', '.join(odd)} differs from the deck font {dom}.",
                                     [e.id], e.bbox, True, dict(op="font", id=e.id, font=dom)))
    # titles: size and position vs the deck's most common
    ts = [(sl, e) for sl in deck.slides for e in sl.els if e.is_title and e.kind == "text"]
    if len(ts) >= 3:
        sizes = Counter(e.size for _, e in ts)
        dom_size = sizes.most_common(1)[0][0]
        pos = Counter((round(e.left), round(e.top)) for _, e in ts)
        dom_pos = pos.most_common(1)[0][0]
        for sl, e in ts:
            # a title already shrunk to fit its length is a Titles matter, not inconsistency
            shrunk = e.size < dom_size and measure(e, cfg, dom_size)["lines"] > cfg.title_max_lines
            if e.size != dom_size and not shrunk:
                out.append(Issue(sl.no, "Consistency",
                                 f"Title is {e.size:g}pt; other slides use {dom_size:g}pt.",
                                 [e.id], e.bbox, True, dict(op="title_size", id=e.id, size=dom_size)))
            if max(abs(e.left - dom_pos[0]), abs(e.top - dom_pos[1])) > cfg.pos_tol:
                out.append(Issue(sl.no, "Consistency",
                                 f"Title position ({e.left:.0f},{e.top:.0f}) differs from other slides {dom_pos}.",
                                 [e.id], e.bbox, True, dict(op="title_pos", id=e.id, pos=dom_pos)))
    # page numbers: present on all or none (title slide exempt)
    body = deck.slides[1:] if len(deck.slides) > 2 else deck.slides
    have = [sl for sl in body if any(e.is_number for e in sl.els)]
    if have and len(have) < len(body):
        ref = next(e for e in have[0].els if e.is_number)
        for sl in body:
            if sl not in have:
                out.append(Issue(sl.no, "Consistency", "Slide has no page number but other slides do.",
                                 [], (ref.left, ref.top, ref.right, ref.bottom), True,
                                 dict(op="add_number", ref=(ref.left, ref.top, ref.width, ref.height))))
    return out


ALL = [alignment, spacing, text_fit, lines_arrows, titles, readability, leftovers, consistency]


def run_all(deck, cfg):
    issues = []
    for fn in ALL:
        issues.extend(fn(deck, cfg))
    return sorted(issues, key=lambda i: (i.slide, i.category))
