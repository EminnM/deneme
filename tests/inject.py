"""Inject known defects into a real deck; returns ground truth [(slide, category)]."""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE


def body_text(slide):
    return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()
            and not (s.is_placeholder and s.placeholder_format.type in (1, 3))]  # not title/center title


def inject(src, dst):
    prs = Presentation(src)
    S = list(prs.slides)
    truth = []

    def pick(n):
        t = body_text(S[n - 1])
        return t[0] if t else None

    # 2: nudge a shape 5pt right -> near-miss alignment needs a sibling; create sibling aligned then offset
    t = pick(2)
    if t:
        a = S[1].shapes.add_textbox(t.left, t.top + Inches(3), Inches(2), Inches(0.5)); a.text_frame.text = "Sibling A"
        b = S[1].shapes.add_textbox(t.left + Pt(5), t.top + Inches(3.8), Inches(2), Inches(0.5)); b.text_frame.text = "Sibling B"
        for x in (a, b):
            x.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        truth.append((2, "Alignment"))
    # 3: tiny text
    t = pick(3)
    if t:
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size = Pt(7)
        truth.append((3, "Readability"))
    # 4: leftover grey box + crooked line
    g = S[3].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), Inches(5), Inches(1.5), Inches(0.8))
    g.fill.solid(); g.fill.fore_color.rgb = RGBColor(0x60, 0x60, 0x60)
    S[3].shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(6.5), Inches(5), Inches(6.62))
    truth += [(4, "Leftover graphics"), (4, "Lines and arrows")]
    # 5: line through text
    t = pick(5)
    if t:
        y = t.top + t.height // 2
        S[4].shapes.add_connector(MSO_CONNECTOR.STRAIGHT, t.left - Inches(0.3), y, t.left + t.width + Inches(0.3), y)
        truth.append((5, "Lines and arrows"))
    # 6: three boxes with uneven gaps
    for i, x in enumerate((0.5, 3.0, 8.0)):
        bx = S[5].shapes.add_textbox(Inches(x), Inches(6.2), Inches(2), Inches(0.5)); bx.text_frame.text = f"Step {i+1}"
        bx.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    truth.append((6, "Spacing"))
    # 7: wrong font
    t = pick(7)
    if t:
        for r in t.text_frame.paragraphs[0].runs:
            r.font.name = "Comic Sans MS"
        truth.append((7, "Consistency"))
    # 8: overflowing fixed box
    t = pick(8)
    if t:
        t.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        t.text_frame.word_wrap = True
        t.text_frame.text = "This paragraph is much too long for the tiny box that holds it so it spills out of the box " * 2
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        t.height = Inches(0.4)
        truth.append((8, "Text fit"))
    # 9: very long title
    ti = S[8].shapes.title
    if ti is not None:
        ti.text_frame.text = "An extremely long title that nobody bothered to shorten and that keeps on going " * 2
        truth.append((9, "Titles"))
    prs.save(dst)
    return truth


if __name__ == "__main__":
    print(inject(sys.argv[1], sys.argv[2]))
