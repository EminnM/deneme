import os, sys, tempfile, unittest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Pt, Inches
from deckqa import model, checks, fixer
from deckqa.config import Config
import make_dense


class Dense(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.clean = os.path.join(cls.tmp, "clean.pptx")
        cls.bad = os.path.join(cls.tmp, "bad.pptx")
        make_dense.build(cls.clean, defects=False)
        make_dense.build(cls.bad, defects=True)
        cls.cfg = Config()

    def run_checks(self, path):
        deck = model.load(path, self.cfg)
        return deck, checks.run_all(deck, self.cfg)

    def test_clean_dense_slide_has_no_false_positives(self):
        _, issues = self.run_checks(self.clean)
        self.assertEqual([i.message for i in issues], [])

    def test_each_seeded_defect_found(self):
        _, issues = self.run_checks(self.bad)
        msgs = " | ".join(i.message for i in issues)
        self.assertIn("off the shared top line (+22pt)", msgs)          # staggered tab (peer outlier, across groups)
        self.assertIn("sits +5,-4pt", msgs)                              # icon off vs sibling panels
        self.assertIn("sits +3,+0pt", msgs)                              # icon nearly centred
        self.assertIn("overlaps", msgs)                                  # text over text
        self.assertIn("wider than its panel", msgs)                      # text spilling out of its panel
        self.assertTrue(any(i.category == "Titles" and "above" in i.message for i in issues))  # title gap

    def test_fix_moves_shapes_inside_groups(self):
        deck, issues = self.run_checks(self.bad)
        first = len(issues)
        out = os.path.join(self.tmp, "fixed.pptx")
        pending = issues
        for _ in range(4):          # same multi-pass loop the CLI uses
            if not fixer.apply(deck, pending, self.cfg):
                break
            deck.prs.save(out)
            deck, after = self.run_checks(out)
            pending = [i for i in after if i.fixable]
            if not pending:
                break
        deck.prs.save(out)
        _, after = self.run_checks(out)
        self.assertEqual([i.message for i in after if i.fixable], [])
        self.assertLess(len(after), first)

    def test_group_transform_is_applied(self):
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        g = s.shapes.add_group_shape()
        g.shapes.add_textbox(Pt(100), Pt(100), Pt(50), Pt(20)).text_frame.text = "hello world"
        x = g._element.xpath("./p:grpSpPr/a:xfrm")[0]
        ext = x.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ext")
        ext.set("cx", str(int(int(ext.get("cx")) * 2)))   # group drawn twice as wide as its child space
        ext.set("cy", str(int(int(ext.get("cy")) * 3)))
        p = os.path.join(self.tmp, "g.pptx")
        prs.save(p)
        deck = model.load(p, self.cfg)
        child = [e for e in deck.slides[0].leaves if e.kind == "text"][0]
        self.assertAlmostEqual(child.left, 100, places=0)        # group offset == child offset
        self.assertAlmostEqual(child.width, 100, places=0)       # 50 * 2
        self.assertAlmostEqual(child.height, 60, places=0)       # 20 * 3
        # moving it in absolute space round-trips through the group transform
        fixer.set_abs(child, left=250, top=40)
        deck.prs.save(p)
        child2 = [e for e in model.load(p, self.cfg).slides[0].leaves if e.kind == "text"][0]
        self.assertAlmostEqual(child2.left, 250, places=0)
        self.assertAlmostEqual(child2.top, 40, places=0)


if __name__ == "__main__":
    unittest.main()
