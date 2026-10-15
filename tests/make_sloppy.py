"""A realistic-looking, deliberately sloppy 8-slide deck. build() returns ground truth:
list of (slide, category, note) for every defect planted."""
import io
from PIL import Image
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

NAVY, TEAL, GREY = RGBColor(0x1B, 0x2A, 0x5E), RGBColor(0x2E, 0x7D, 0x8C), RGBColor(0x80, 0x80, 0x80)


def icon(color):
    im = Image.new("RGBA", (64, 64), (255, 255, 255, 0))
    for x in range(10, 54):
        for y in range(10, 54):
            if (x - 32) ** 2 + (y - 32) ** 2 < 22 ** 2:
                im.putpixel((x, y), color + (255,))
    b = io.BytesIO(); im.save(b, "PNG"); b.seek(0)
    return b


def tb(shapes, l, t, w, h, text, size=14, bold=False, color=None, font="Calibri", align=None, fixed=True, wrap=True):
    b = shapes.add_textbox(Pt(l), Pt(t), Pt(w), Pt(h))
    tf = b.text_frame
    tf.word_wrap = wrap
    if fixed:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.text = text
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    r = p.runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    if color is not None:
        r.font.color.rgb = color
    return b


def rect(shapes, l, t, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    r = shapes.add_shape(shape, Pt(l), Pt(t), Pt(w), Pt(h))
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
    return r


def build(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    L_TITLE, L_CONTENT, L_TONLY = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]
    truth = []

    def content_slide(title, layout=None, size=28, pos=(36, 24), number=True):
        s = prs.slides.add_slide(layout or L_TONLY)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Pt(pos[0]), Pt(pos[1]), Pt(888), Pt(60)
        t.text = title
        r = t.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(size), True, "Calibri"
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        if number:
            tb(s.shapes, 880, 505, 44, 22, str(len(prs.slides)), 11, align=PP_ALIGN.RIGHT)
        return s

    # ---------------------------------------------------------------- 1 title
    s = prs.slides.add_slide(L_TITLE)
    s.shapes.title.text = "Altheria Mental Health Strategy"
    s.placeholders[1].text = "Board discussion document - Q3"
    g = rect(s.shapes, 820, 440, 90, 50, GREY)          # forgotten grey box
    truth.append((1, "Leftover graphics", "grey box"))

    # ---------------------------------------------------------------- 2 agenda
    s = content_slide("Agenda")
    items = ["Context and ambition", "Where the system struggles", "Evidence and benchmarks",
             "Options and recommendation", "Roadmap and next steps"]
    for i, it in enumerate(items):
        y = 120 + i * 62 + (24 if i == 4 else 0)              # last row gap too big
        cx = 60 + (14 if i == 3 else 0)                       # circle pushed right
        c = rect(s.shapes, cx, y, 40, 40, TEAL, MSO_SHAPE.OVAL)
        c.text_frame.text = str(i + 1)
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        ty = y + 6 + (2 if i == 2 else 0)                     # one text box 2pt low
        tb(s.shapes, 120, ty, 500, 28, it, 20)
    truth += [(2, "Alignment", "circle pushed right"), (2, "Alignment", "text 2pt low"), (2, "Spacing", "uneven gap")]

    # ---------------------------------------------------------------- 3 three columns
    s = content_slide("Three structural barriers drive the gap")
    bodies = ["Few trained specialists outside the capital; long waits for first appointment and almost no follow-up care.",
              "Funding flows to hospital beds instead of community services, so patients are admitted late and stay long. "
              "Referral pathways are unclear, discharge planning is rare, and families carry most of the burden at home "
              "without support, training or respite options of any kind across regions. Private providers fill part of "
              "the gap but are unaffordable for most households, and insurance coverage for outpatient therapy is "
              "limited to a handful of sessions per year, which pushes people back to emergency departments. "
              "Rural districts report waiting lists of more than six months for any psychiatric appointment.",
              "Stigma keeps people from seeking help early; awareness campaigns are fragmented."]
    for i in range(3):
        x = 36 + i * 300
        g = s.shapes.add_group_shape()
        rect(g.shapes, x, 130, 280, 340, RGBColor(0xE8, 0xEE, 0xF2))
        rect(g.shapes, x, 100 + (18 if i == 1 else 0), 280, 30, NAVY)       # header bar
        ic_dx, ic_dy = 120 + (4 if i == 2 else 0), 134 + (3 if i == 2 else 0)
        g.shapes.add_picture(icon((46, 125, 140)), Pt(x + ic_dx - 36 + 5), Pt(ic_dy), Pt(40), Pt(40))
        font = "Comic Sans MS" if i == 2 else "Calibri"
        tb(g.shapes, x + 10, 190, 260, 250, bodies[i], 14, font=font)
    tb(s.shapes, 36, 478, 600, 14, "Source: Ministry of Health annual report; team analysis", 8)
    truth += [(3, "Alignment", "header bar staggered"), (3, "Alignment", "icon off"), (3, "Text fit", "col 2 too long"),
              (3, "Consistency", "comic sans"), (3, "Readability", "8pt footnote")]

    # ---------------------------------------------------------------- 4 table + chart
    s = content_slide("Mental health spending is far below peer benchmarks and the gap with comparable countries is widening every single year since 2018, which puts the whole reform agenda at risk", size=28)
    rows, cols = 5, 3
    gt = s.shapes.add_table(rows, cols, Pt(36), Pt(120), Pt(420), Pt(200))
    data = [["Country", "Spend per capita", "Beds /100k"], ["Altheria", "$12", "18"], ["Peer A", "$46", "41"],
            ["Peer B", "$51", "47"], ["Peer C", "$39", "35"]]
    for r in range(rows):
        for c in range(cols):
            gt.table.cell(r, c).text = data[r][c]
    cd = CategoryChartData(); cd.categories = ["2018", "2019", "2020", "2021"]
    cd.add_series("Altheria", (10, 11, 12, 12)); cd.add_series("Peers", (38, 41, 44, 46))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(420), Pt(140), Pt(480), Pt(300), cd)   # overlaps table
    tb(s.shapes, 36, 470, 500, 14, "Source: WHO Mental Health Atlas", 7)
    truth += [(4, "Titles", "3-line title"), (4, "Spacing", "chart overlaps table"), (4, "Readability", "7pt source")]

    # ---------------------------------------------------------------- 5 process
    s = content_slide("Four-step delivery approach")
    labels = ["Diagnose", "Design", "Implementation", "Scale"]
    for i, lb in enumerate(labels):
        x = 36 + i * 225
        ch = rect(s.shapes, x, 150 - (10 if i == 2 else 0), 205, 70, TEAL, MSO_SHAPE.CHEVRON)
        ch.text_frame.text = lb
        ch.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        tb(s.shapes, x + 20, 240, 150 if i != 2 else 60, 80,
           "Assess current state" if i != 2 else "Implementation of regional rollout", 14)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(36), Pt(262), Pt(900), Pt(262))     # runs through text
    c2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(100), Pt(400), Pt(500), Pt(412))    # crooked
    blob = rect(s.shapes, 650, 380, 70, 70, RGBColor(0x30, 0x60, 0xD0), MSO_SHAPE.OVAL)       # stray blob
    truth += [(5, "Alignment", "chevron raised"), (5, "Lines and arrows", "line through text"),
              (5, "Lines and arrows", "crooked"), (5, "Leftover graphics", "blob"), (5, "Text fit", "narrow box")]

    # ---------------------------------------------------------------- 6 text heavy
    s = content_slide("Detailed findings", layout=L_CONTENT, size=32, pos=(36, 40))
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Pt(36), Pt(110), Pt(888), Pt(320)
    tf = body.text_frame
    pts = [f"Finding {i + 1}: workforce, funding and governance gaps compound each other across regions" for i in range(13)]
    tf.text = pts[0]
    for t_ in pts[1:]:
        tf.add_paragraph().text = t_
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(22)
    tb(s.shapes, 880, 505, 44, 22, "6", 11, align=PP_ALIGN.RIGHT)
    truth += [(6, "Text fit", "bullets overflow"), (6, "Consistency", "title size/pos")]

    # ---------------------------------------------------------------- 7 card grid (no page number)
    s = content_slide("Six levers for system change", layout=L_CONTENT, number=False)
    for i in range(6):
        r_, c_ = divmod(i, 3)
        x = 36 + c_ * 300 + (12 if i == 4 else 0)
        y = 120 + r_ * 180
        w = 280 + (6 if i == 1 else 0)
        g = s.shapes.add_group_shape()
        rect(g.shapes, x, y, w, 160, RGBColor(0xEE, 0xF1, 0xF5))
        g.shapes.add_picture(icon((30, 70, 110)), Pt(x + 12), Pt(y + 12), Pt(36), Pt(36))
        tb(g.shapes, x + 56, y + 14, w - 70, 130, f"Lever {i + 1}: integrate care pathways between primary and secondary care", 14)
    truth += [(7, "Alignment", "card width 6pt"), (7, "Alignment", "card shifted 12pt"),
              (7, "Leftover graphics", "empty body placeholder"), (7, "Consistency", "no page number")]

    # ---------------------------------------------------------------- 8 next steps
    s = content_slide("Next steps and decisions requested")
    tb(s.shapes, 36, 66, 300, 100, "Decision 1: approve the funding envelope for year one", 16)   # box tucked under the title
    tb(s.shapes, 39, 220, 300, 100, "Decision 2: appoint a programme director", 16)
    tb(s.shapes, 41, 350, 300, 100, "Decision 3: confirm the pilot regions", 16)
    tb(s.shapes, 400, 220, 500, 14, "Timeline subject to ministry sign-off", 8)
    rect(s.shapes, 700, 380, 120, 60, GREY)
    truth += [(8, "Alignment", "lefts 36/39/41"), (8, "Readability", "8pt"), (8, "Leftover graphics", "grey box"),
              (8, "Titles", "box touching title")]

    prs.save(path)
    return truth


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1]))
