"""Builds a small deck with one deliberate defect per checklist category."""
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR


def build(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank, title_only = prs.slide_layouts[6], prs.slide_layouts[5]

    def slide(title, size=32, number=True, pos=(0.6, 0.4)):
        s = prs.slides.add_slide(title_only)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(pos[0]), Inches(pos[1]), Inches(12), Inches(1.0)
        t.text = title
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        t.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        if number:
            n = s.shapes.add_textbox(Inches(12.3), Inches(6.9), Inches(0.6), Inches(0.4))
            n.text_frame.text = str(len(prs.slides))
            n.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        return s

    def box(s, l, t, w, h, text, size=16, font="Arial"):
        b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        b.text_frame.word_wrap = True
        b.text_frame.auto_size = MSO_AUTO_SIZE.NONE  # fixed-size box, like a designed card
        b.text_frame.text = text
        r = b.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.name = Pt(size), font
        return b

    # 1 clean-ish
    s = slide("Overview")
    box(s, 0.6, 1.8, 5.5, 1.2, "Clean left block", 18)
    box(s, 6.8, 1.8, 5.5, 1.2, "Clean right block", 18)
    # 2 alignment + spacing
    s = slide("Alignment and spacing")
    box(s, 0.6, 1.8, 3.5, 1.0, "Card A", 18)
    box(s, 4.7, 1.85, 3.5, 1.0, "Card B", 18)    # top off by 0.05in (3.6pt)
    box(s, 9.6, 1.8, 3.0, 1.0, "Card C", 18)     # uneven gap
    # 3 text fit + readability
    s = slide("Text fit and readability")
    box(s, 0.6, 1.8, 3.0, 0.5, "This text is far too long for the little box it was put into so it spills out", 18)
    box(s, 0.6, 5.0, 6.0, 0.5, "Tiny footnote", 7)
    # 4 lines and arrows
    s = slide("Lines")
    box(s, 4.0, 3.0, 4.0, 1.0, "Important text", 18)
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3.0), Inches(3.5), Inches(9.0), Inches(3.5))
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.0), Inches(5.5), Inches(5.0), Inches(5.62))
    # 5 titles
    s = slide("A very long title that clearly goes on and on across several lines because nobody shortened it for the slide and it keeps going with even more words that make it far too long to fit in two lines", 32)
    box(s, 0.6, 1.3, 6.0, 1.0, "Body that the title crashes into", 18)
    # 6 leftover graphics
    s = slide("Leftover shapes")
    g = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(2), Inches(1))
    g.fill.solid(); g.fill.fore_color.rgb = RGBColor(0x80, 0x80, 0x80)
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6), Inches(3), Inches(1), Inches(1))
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xE0, 0x30, 0x90)
    # 7 consistency: other font, other title size/position, no page number
    s = slide("Inconsistent slide", size=26, number=False, pos=(1.4, 0.7))
    box(s, 0.6, 1.8, 6.0, 1.0, "Comic Sans text in an Arial deck", 18, font="Comic Sans MS")
    prs.save(path)


if __name__ == "__main__":
    import sys
    build(sys.argv[1] if len(sys.argv) > 1 else "sample_defects.pptx")
